import os, io, re, sqlite3, random, logging, json, asyncio, zipfile
from datetime import datetime, timedelta, timezone

from dotenv import load_dotenv
from telegram import (
    Update, ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardMarkup,
    InlineKeyboardButton
)
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    ContextTypes, filters
)
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))
BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_LOGIN = os.getenv("ADMIN_LOGIN", "admin").strip()
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "").strip()
# Railway PostgreSQL ulanganda DATABASE_URL avtomatik beriladi. Bo'lmasa lokal SQLite fayl ishlatiladi.
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
USE_PG = DATABASE_URL.startswith(("postgres://", "postgresql://"))
DB_PATH = os.getenv("DB_PATH", "icu_test.db")
if not os.path.isabs(DB_PATH):
    DB_PATH = os.path.join(BASE_DIR, DB_PATH)
if USE_PG:
    import psycopg
    from psycopg.rows import dict_row
ACADEMY = "AL-AZIZ ACADEMY"
# Toshkent vaqti (UTC+5). Server qayerda bo'lishidan qat'i nazar, vaqt shu bo'yicha hisoblanadi.
TZ = timezone(timedelta(hours=5))
TOTAL_Q = 30
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

# Each participant gets one random variant. These are starter IQ/logic questions;
# replace/edit them in QUESTION_BANK if your ICU curriculum uses specific questions.
QUESTION_BANK = {
  1: [
    ("2 + 3 = ?", ["4","5","6","7"], 1),
    ("Qaysi biri meva?", ["Sabzi","Olma","Kartoshka","Piyoz"], 1),
    ("1, 2, 3, __", ["4","5","6","7"], 0),
    ("Haftada nechta kun bor?", ["5","6","7","8"], 2),
    ("Qaysi shakl dumaloq?", ["Kvadrat","Uchburchak","Doira","To‘rtburchak"], 2),
    ("10 dan oldingi son?", ["8","9","11","12"], 1),
    ("Qaysi biri hayvon?", ["Stol","Mushuk","Kitob","Qalam"], 1),
    ("4 + 4 = ?", ["6","7","8","9"], 2),
    ("Qaysi rang osmonning kunduzgi rangiga ko‘proq mos?", ["Ko‘k","Qora","Jigarrang","Pushti"], 0),
    ("5, 6, 7, __", ["8","9","10","11"], 0),
    ("Qaysi biri maktab buyumi?", ["Qalam","Non","Olma","Sut"], 0),
    ("9 - 3 = ?", ["5","6","7","8"], 1),
    ("Bir qo‘lda odatda nechta barmoq bor?", ["3","4","5","6"], 2),
    ("Qaysi biri transport?", ["Avtobus","Daraxt","Daftar","Gul"], 0),
    ("3 + 6 = ?", ["7","8","9","10"], 2),
    ("Qishdan keyin qaysi fasl keladi?", ["Kuz","Bahor","Yoz","Qish"], 1),
    ("Eng kichik son qaysi?", ["7","2","5","9"], 1),
    ("12 - 2 = ?", ["8","9","10","11"], 2),
    ("Qaysi biri ichimlik?", ["Suv","Non","Guruch","Tuz"], 0),
    ("2, 4, 6, __", ["7","8","9","10"], 1),
    ("Bir haftada nechta kun dam olish kuni odatda bo‘ladi?", ["1","2","4","7"], 1),
    ("Qaysi hayvon uchadi?", ["Baliq","Qush","Mushuk","Ot"], 1),
    ("7 + 1 = ?", ["6","7","8","9"], 2),
    ("Qaysi son 10 dan katta?", ["8","9","10","12"], 3),
    ("Qaysi biri sabzavot?", ["Uzum","Banan","Sabzi","Nok"], 2),
    ("5 - 2 = ?", ["2","3","4","5"], 1),
    ("Qaysi shaklning 3 tomoni bor?", ["Doira","Kvadrat","Uchburchak","Oval"], 2),
    ("1, 3, 5, __", ["6","7","8","9"], 1),
    ("Qaysi biri kiyim?", ["Ko‘ylak","Likopcha","Qoshiq","Ruchka"], 0),
    ("6 + 2 = ?", ["7","8","9","10"], 1),
  ],
  2: [
    ("15 + 7 = ?", ["20","21","22","23"], 2),
    ("Qatorni davom ettiring: 2, 4, 8, 16, __", ["18","24","30","32"], 3),
    ("Bir soat necha daqiqa?", ["30","45","60","100"], 2),
    ("36 ÷ 6 = ?", ["5","6","7","8"], 1),
    ("Qaysi so‘z ortiqcha: olma, nok, sabzi, uzum?", ["Olma","Nok","Sabzi","Uzum"], 2),
    ("25 - 9 = ?", ["14","15","16","17"], 2),
    ("Uchta 5 ning yig‘indisi?", ["10","15","20","25"], 1),
    ("Yilning nechta oyi bor?", ["10","11","12","13"], 2),
    ("9 × 4 = ?", ["32","34","36","40"], 2),
    ("100 ning yarmi?", ["25","40","50","60"], 2),
    ("Qaysi son juft?", ["13","17","21","24"], 3),
    ("48 + 12 = ?", ["50","58","60","62"], 2),
    ("Ketma-ketlik: 5, 10, 15, __", ["18","20","25","30"], 1),
    ("Kvadratning nechta tomoni bor?", ["3","4","5","6"], 1),
    ("7 × 7 = ?", ["42","47","49","56"], 2),
    ("Bir metr necha santimetr?", ["10","50","100","1000"], 2),
    ("64 ÷ 8 = ?", ["6","7","8","9"], 2),
    ("Qaysi biri materik?", ["Afrika","Tinch okeani","Nil","Amazonka"], 0),
    ("90 - 35 = ?", ["45","50","55","65"], 2),
    ("Agar bugun seshanba bo‘lsa, ertaga qaysi kun?", ["Dushanba","Chorshanba","Payshanba","Juma"], 1),
    ("3 × (4 + 2) = ?", ["14","18","20","24"], 1),
    ("Eng katta son qaysi?", ["0.7","0.07","0.77","0.707"], 2),
    ("Bir kilogramm necha gramm?", ["100","500","1000","10000"], 2),
    ("144 ÷ 12 = ?", ["10","11","12","14"], 2),
    ("Qaysi son 3 ga bo‘linadi?", ["14","16","18","20"], 2),
    ("20 ning 25 foizi?", ["4","5","10","15"], 1),
    ("Qaysi so‘z fe’l?", ["Chiroyli","Yugurmoq","Kitob","Tez"], 1),
    ("2, 5, 8, 11, __", ["12","13","14","15"], 2),
    ("To‘g‘ri to‘rtburchakning qarama-qarshi tomonlari...", ["Har doim turli","Teng","Doira","Yo‘q"], 1),
    ("81 ning kvadrat ildizi?", ["7","8","9","10"], 2),
  ],
  3: [
    ("Agar 3x = 27 bo‘lsa, x = ?", ["6","7","8","9"], 3),
    ("Ketma-ketlik: 1, 1, 2, 3, 5, __", ["6","7","8","9"], 2),
    ("200 ning 15 foizi?", ["15","20","30","35"], 2),
    ("Tenglama: x + 14 = 31. x = ?", ["15","16","17","18"], 2),
    ("Qaysi son tub son?", ["21","27","29","33"], 2),
    ("2³ × 3 = ?", ["12","18","24","36"], 2),
    ("Uchburchak ichki burchaklari yig‘indisi?", ["90°","180°","270°","360°"], 1),
    ("0.6 kasr ko‘rinishida?", ["3/5","2/5","1/6","6/100"], 0),
    ("Agar barcha A lar B bo‘lsa, va X — A bo‘lsa, X nima?", ["B ham","B emas","Aniqlab bo‘lmaydi","A emas"], 0),
    ("12² = ?", ["124","122","144","164"], 2),
    ("80 ning 3/4 qismi?", ["20","40","60","70"], 2),
    ("Ketma-ketlik: 4, 9, 16, 25, __", ["30","32","36","49"], 2),
    ("5x - 10 = 20. x = ?", ["4","5","6","8"], 2),
    ("1, 4, 9, 16, 25, __", ["30","32","36","49"], 2),
    ("Agar 2:5 = x:20 bo‘lsa, x = ?", ["6","8","10","12"], 1),
    ("Qaysi biri mantiqan to‘g‘ri: barcha baliqlar suvda yashaydi; forel baliq. Demak...", ["Forel suvda yashaydi","Forel qush","Barcha suv jonivori baliq","Xulosa yo‘q"], 0),
    ("360° ning 1/8 qismi?", ["30°","45°","60°","90°"], 1),
    ("Bir sonning 40 foizi 28. Son nechchi?", ["56","60","70","80"], 2),
    ("2x + 7 = 19. x = ?", ["5","6","7","8"], 1),
    ("Qaysi kasr eng katta?", ["2/3","3/5","5/8","1/2"], 0),
    ("To‘g‘ri burchak necha gradus?", ["45°","90°","120°","180°"], 1),
    ("Agar bugun oyning 10-kuni bo‘lsa, 17 kundan keyin?", ["25","26","27","28"], 2),
    ("7² - 5² = ?", ["12","24","25","49"], 1),
    ("Aylana uzunligi formulasi?", ["πr²","2πr","a²","ab"], 1),
    ("Agar 4 ishchi 6 kunda ishni bajarsa, bir xil tezlikda 8 ishchi necha kunda?", ["2","3","4","12"], 1),
    ("3, 6, 12, 24, __", ["30","36","42","48"], 3),
    ("1/3 + 1/6 = ?", ["1/2","2/9","1/9","1"], 0),
    ("Agar P > Q va Q > R bo‘lsa, qaysi biri to‘g‘ri?", ["R > P","P > R","Q < R","P = R"], 1),
    ("150 ning 20 foizi?", ["15","20","30","50"], 2),
    ("√225 = ?", ["12","13","14","15"], 3),
  ],
}
# Normalize exactly 30 questions per variant.
assert all(len(v) == 30 for v in QUESTION_BANK.values())

# ---------------------------------------------------------------- baza

_conn = None

def _connect():
    global _conn
    if USE_PG:
        if _conn is None or _conn.closed:
            _conn = psycopg.connect(DATABASE_URL, autocommit=True, row_factory=dict_row)
    elif _conn is None:
        _conn = sqlite3.connect(DB_PATH, isolation_level=None)
        _conn.row_factory = sqlite3.Row
    return _conn

def sql(query, params=(), fetch=None):
    """Bitta SQL so'rov. fetch: None, "one" yoki "all"."""
    global _conn
    if USE_PG:
        query = query.replace("?", "%s")
    for attempt in range(2):
        try:
            cur = _connect().execute(query, params)
            if fetch == "one": return cur.fetchone()
            if fetch == "all": return cur.fetchall()
            return None
        except Exception as e:
            # Ulanish uzilsa (masalan, baza qayta ishga tushsa) bir marta qayta ulanamiz
            if USE_PG and attempt == 0 and isinstance(e, psycopg.OperationalError):
                log.warning("Bazaga qayta ulanish: %s", e)
                try: _conn.close()
                except Exception: pass
                _conn = None
                continue
            raise

def init_db():
    sql("""CREATE TABLE IF NOT EXISTS users (
      tg_id BIGINT PRIMARY KEY, username TEXT, full_name TEXT, phone TEXT,
      branch TEXT, grade INTEGER, variant INTEGER, status TEXT DEFAULT 'registered',
      score INTEGER, total INTEGER, started_at TEXT, finished_at TEXT,
      registered_at TEXT NOT NULL, question_order TEXT, answers TEXT, school TEXT
    )""")
    if USE_PG:
        sql("ALTER TABLE users ADD COLUMN IF NOT EXISTS school TEXT")
    else:
        cols = {r["name"] for r in sql("PRAGMA table_info(users)", fetch="all")}
        if "school" not in cols:
            sql("ALTER TABLE users ADD COLUMN school TEXT")
    sql("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)")
    sql("CREATE TABLE IF NOT EXISTS admins (tg_id BIGINT PRIMARY KEY)")

def now():
    return datetime.now(TZ)

def now_str():
    return now().strftime("%Y-%m-%d %H:%M:%S")

def get_user(tg_id):
    return sql("SELECT * FROM users WHERE tg_id=?", (tg_id,), "one")

def set_fields(tg_id, **fields):
    if not fields: return
    keys = list(fields)
    sql("UPDATE users SET " + ", ".join(f"{k}=?" for k in keys) + " WHERE tg_id=?",
        [fields[k] for k in keys] + [tg_id])

def upsert_user(tg_id, username):
    sql("""INSERT INTO users(tg_id, username, status, registered_at)
      VALUES(?,?,?,?) ON CONFLICT(tg_id) DO UPDATE SET username=excluded.username""",
      (tg_id, username or "", "new", now_str()))

def get_setting(key):
    r = sql("SELECT value FROM settings WHERE key=?", (key,), "one")
    return r["value"] if r else None

def set_setting(key, value):
    if value is None:
        sql("DELETE FROM settings WHERE key=?", (key,))
    else:
        sql("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))

def get_start_time():
    v = get_setting("test_start")
    if not v: return None
    try: return datetime.strptime(v, "%Y-%m-%d %H:%M").replace(tzinfo=TZ)
    except ValueError: return None

def fmt_dt(dt):
    return dt.strftime("%d.%m.%Y") + " kuni soat " + dt.strftime("%H:%M")

def is_admin(tg_id):
    return sql("SELECT 1 AS x FROM admins WHERE tg_id=?", (tg_id,), "one") is not None

def admin_ids():
    return [r["tg_id"] for r in sql("SELECT tg_id FROM admins", fetch="all")]

def json_load(v):
    try: return json.loads(v or "[]")
    except Exception: return []

def feedback_for(percent):
    if percent < 50:
        return "Ko‘proq kitob o‘qing va yana mashq qiling!"
    if percent < 70:
        return "Harakatingiz yaxshi! Mashq qilishda davom eting."
    if percent < 80:
        return "Yaxshi natija! Yana biroz harakat qiling."
    if percent <= 90:
        return "Ajoyib natija! Shunday davom eting."
    return "Juda zo‘r natija! Siz haqiqiy bilimdonsiz!"

# ---------------------------------------------------------------- klaviaturalar

BTN_REGISTER = "📝 Ro‘yxatdan o‘tish"
BTN_ADMIN = "🔐 Admin panelga kirish"
BTN_START = "🧠 IQ testni boshlash"
BTN_CONTINUE = "▶️ Davom etish"
BTN_FINISH = "🏁 Testni yakunlash"
BTN_MENU = "⬅️ Bosh menyu"
BTN_STATS = "📊 Statistika"
BTN_USERS = "👥 Foydalanuvchilar"
BTN_EXCEL = "📥 Excel yuklab olish"
BTN_CERTS = "🏅 Sertifikatlarni yuklab olish"
BTN_SET_TIME = "⏰ Test vaqtini belgilash"
BTN_CLEAR_TIME = "🗑 Test vaqtini bekor qilish"
BTN_LOGOUT = "🚪 Admin chiqish"

def kb(rows, **kw):
    return ReplyKeyboardMarkup(rows, resize_keyboard=True, **kw)

def main_keyboard(u=None):
    status = u["status"] if u else None
    if status == "in_progress":
        return kb([[BTN_CONTINUE], [BTN_FINISH]])
    if status == "registered":
        return kb([[BTN_START], [BTN_REGISTER, BTN_ADMIN]])
    if status == "completed":
        return kb([[BTN_ADMIN]])
    return kb([[BTN_REGISTER], [BTN_ADMIN]])

def back_keyboard():
    return kb([[BTN_MENU]])

def phone_keyboard():
    return kb([[KeyboardButton("📱 Telefon raqamni yuborish", request_contact=True)], [BTN_MENU]], one_time_keyboard=True)

def grade_keyboard():
    nums = [str(i) for i in range(1, 12)]
    rows = [nums[i:i+4] for i in range(0, len(nums), 4)]
    rows.append([BTN_MENU])
    return kb(rows, one_time_keyboard=True)

def admin_keyboard():
    return kb([
        [BTN_SET_TIME, BTN_CLEAR_TIME],
        [BTN_STATS, BTN_USERS],
        [BTN_EXCEL, BTN_CERTS],
        [BTN_LOGOUT, BTN_MENU],
    ])

# ---------------------------------------------------------------- sertifikat

FONT_DIR = os.path.join(BASE_DIR, "fonts")

def _register_fonts():
    weights = {"Light": "300Light", "Regular": "400Regular", "Medium": "500Medium",
               "SemiBold": "600SemiBold", "Bold": "700Bold", "ExtraBold": "800ExtraBold"}
    fallback = {"Light": "Helvetica", "Regular": "Helvetica", "Medium": "Helvetica",
                "SemiBold": "Helvetica-Bold", "Bold": "Helvetica-Bold", "ExtraBold": "Helvetica-Bold"}
    fonts = {}
    for key, w in weights.items():
        path = os.path.join(FONT_DIR, f"Montserrat_{w}.ttf")
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("Mont" + key, path))
            fonts[key] = "Mont" + key
        else:
            log.warning("Shrift topilmadi: %s", path)
            fonts[key] = fallback[key]
    return fonts

F = _register_fonts()
NAVY = HexColor("#0E1F3D")
NAVY_2 = HexColor("#1D3F7A")
GOLD = HexColor("#D4A537")
GOLD_2 = HexColor("#F3D27A")
INK = HexColor("#16233A")
MUTED = HexColor("#7A8699")
CARD = HexColor("#F3F6FB")
WHITE = HexColor("#FFFFFF")

def _fit_size(text, font, max_size, max_width, min_size=10):
    size = max_size
    while size > min_size and pdfmetrics.stringWidth(text, font, size) > max_width:
        size -= 0.5
    return size

def _spaced(c, x, y, text, font, size, spacing, color, anchor="left"):
    """Harflar orasi ochiq yozuv."""
    w = pdfmetrics.stringWidth(text, font, size) + spacing * (len(text) - 1)
    if anchor == "center": x -= w / 2
    t = c.beginText(x, y)
    t.setFont(font, size); t.setCharSpace(spacing); t.setFillColor(color)
    t.textOut(text)
    t.setCharSpace(0)  # keyingi yozuvlarga ta'sir qilmasin
    c.drawText(t)
    return w

def _poly(c, pts, color, alpha=1.0):
    c.saveState()
    c.setFillColor(color); c.setFillAlpha(alpha)
    p = c.beginPath(); p.moveTo(*pts[0])
    for pt in pts[1:]: p.lineTo(*pt)
    p.close(); c.drawPath(p, fill=1, stroke=0)
    c.restoreState()

def _wrap(text, font, size, width):
    lines, cur = [], ""
    for word in text.split():
        trial = (cur + " " + word).strip()
        if pdfmetrics.stringWidth(trial, font, size) <= width: cur = trial
        else:
            if cur: lines.append(cur)
            cur = word
    if cur: lines.append(cur)
    return lines

def make_certificate(name, score, total, grade, school, date_str=None):
    out = io.BytesIO()
    W, H = landscape(A4)
    c = canvas.Canvas(out, pagesize=(W, H))
    c.setTitle(f"Sertifikat - {name}")
    percent = round(score / total * 100) if total else 0
    PW = 268  # chap panel kengligi

    # Oq fon
    c.setFillColor(WHITE); c.rect(0, 0, W, H, fill=1, stroke=0)

    # O‘ng yuqori va past burchakdagi geometrik shakllar
    _poly(c, [(W, H), (W-190, H), (W, H-190)], NAVY, 0.06)
    _poly(c, [(W, H), (W-120, H), (W, H-120)], GOLD, 0.85)
    _poly(c, [(W, H), (W-62, H), (W, H-62)], NAVY, 1)
    _poly(c, [(W, 0), (W-150, 0), (W, 150)], GOLD, 0.12)
    c.saveState(); c.setFillColor(NAVY); c.setFillAlpha(0.10)
    for i in range(6):
        for j in range(4):
            c.circle(W - 40 - i*13, 40 + j*13, 1.6, fill=1, stroke=0)
    c.restoreState()

    # Chap panel: gradient
    c.saveState()
    p = c.beginPath(); p.rect(0, 0, PW, H); c.clipPath(p, stroke=0, fill=0)
    c.linearGradient(0, 0, PW, H, (NAVY, NAVY_2), extend=False)
    for k, a in enumerate((0.05, 0.04, 0.06)):
        off = k * 120
        _poly(c, [(-40 + off, 0), (60 + off, 0), (PW + 40 + off, H), (PW - 60 + off, H)], WHITE, a)
    _poly(c, [(0, 0), (PW, 0), (0, 130)], WHITE, 0.05)
    c.restoreState()
    c.setFillColor(GOLD); c.rect(PW, 0, 5, H, fill=1, stroke=0)

    # Panel ichida: akademiya nomi
    _spaced(c, PW/2, H-70, "AL-AZIZ", F["ExtraBold"], 20, 4, WHITE, "center")
    _spaced(c, PW/2, H-92, "ACADEMY", F["Medium"], 11, 6, GOLD_2, "center")
    c.setStrokeColor(GOLD); c.setLineWidth(1)
    c.line(PW/2 - 24, H-108, PW/2 + 24, H-108)

    # Natija halqasi
    cx, cy, r = PW/2, H/2 - 10, 78
    c.saveState()
    c.setStrokeColor(WHITE); c.setStrokeAlpha(0.14); c.setLineWidth(14)
    c.circle(cx, cy, r, fill=0, stroke=1)
    c.restoreState()
    if percent > 0:
        c.saveState()
        c.setStrokeColor(GOLD); c.setLineWidth(14); c.setLineCap(1)
        arc = c.beginPath()
        arc.arc(cx - r, cy - r, cx + r, cy + r, startAng=90, extent=-3.6 * min(percent, 99.9))
        c.drawPath(arc, stroke=1, fill=0)
        c.restoreState()
    c.setFillColor(WHITE)
    size = 42 if percent < 100 else 36
    c.setFont(F["ExtraBold"], size)
    c.drawCentredString(cx, cy - 8, f"{percent}%")
    _spaced(c, cx, cy - 32, "NATIJA", F["SemiBold"], 9, 3, GOLD_2, "center")

    _spaced(c, PW/2, 70, "IQ TEST", F["Bold"], 12, 4, WHITE, "center")

    # O‘ng tomon: asosiy matn
    X = PW + 62
    RW = W - X - 70
    _spaced(c, X, H-92, "SERTIFIKAT", F["ExtraBold"], 46, 5, NAVY)
    _spaced(c, X + 2, H-118, "IQ TEST ISHTIROKCHISI", F["SemiBold"], 11, 3.5, GOLD)

    c.setFillColor(MUTED); c.setFont(F["Regular"], 12.5)
    c.drawString(X, H-180, "Ushbu sertifikat quyidagi ishtirokchiga taqdim etiladi:")

    safe_name = (name or "Ishtirokchi").strip()[:80]
    nsize = _fit_size(safe_name, F["Bold"], 38, RW, 16)
    c.setFillColor(NAVY); c.setFont(F["Bold"], nsize)
    c.drawString(X, H-230, safe_name)
    c.setFillColor(GOLD); c.rect(X, H-250, 70, 4, fill=1, stroke=0)
    c.setStrokeColor(HexColor("#E3E8F0")); c.setLineWidth(1)
    c.line(X + 78, H-248, X + RW, H-248)

    info = []
    if school: info.append(("Maktab", str(school)))
    if grade: info.append(("Sinf", f"{grade}-sinf"))
    ix = X
    for label, val in info:
        c.setFillColor(MUTED); c.setFont(F["Medium"], 9)
        _spaced(c, ix, H-276, label.upper(), F["SemiBold"], 8.5, 2, MUTED)
        vs = _fit_size(val, F["SemiBold"], 13, RW * 0.62 if label == "Maktab" else 120, 8)
        c.setFillColor(INK); c.setFont(F["SemiBold"], vs)
        c.drawString(ix, H-294, val)
        ix += max(pdfmetrics.stringWidth(val, F["SemiBold"], vs), 60) + 48

    # Izoh kartasi
    fb = feedback_for(percent)
    lines = _wrap(fb, F["SemiBold"], 17, RW - 50)
    ch = 62 + 22 * (len(lines) - 1)
    cyb = H - 340 - ch
    c.setFillColor(CARD); c.roundRect(X, cyb, RW, ch, 10, fill=1, stroke=0)
    c.setFillColor(GOLD); c.roundRect(X, cyb, 6, ch, 3, fill=1, stroke=0)
    _spaced(c, X + 26, cyb + ch - 22, "IZOH", F["Bold"], 9, 3, GOLD)
    c.setFillColor(NAVY); c.setFont(F["SemiBold"], 17)
    for i, ln in enumerate(lines):
        c.drawString(X + 26, cyb + ch - 45 - 22 * i, ln)

    # Imzo
    c.setStrokeColor(HexColor("#C9D1DD")); c.setLineWidth(0.8)
    c.line(X, 92, X + 190, 92)
    c.line(X + RW - 190, 92, X + RW, 92)
    _spaced(c, X, 76, "SANA", F["SemiBold"], 8.5, 2, MUTED)
    _spaced(c, X + RW - 190, 76, "DIREKTOR IMZOSI", F["SemiBold"], 8.5, 2, MUTED)
    c.setFillColor(INK); c.setFont(F["Medium"], 12)
    c.drawString(X, 100, date_str or now().strftime("%d.%m.%Y"))

    c.showPage()
    c.save()
    out.seek(0)
    return out

def safe_filename(name, tg_id):
    base = re.sub(r"[^\w\- ]+", "", name or "", flags=re.UNICODE).strip().replace(" ", "_")
    return f"{base or 'ishtirokchi'}_{tg_id}_sertifikat.pdf"

def certificate_for_row(u):
    score = u["score"] or 0
    total = u["total"] or TOTAL_Q
    date_str = None
    if u["finished_at"]:
        try: date_str = datetime.fromisoformat(u["finished_at"]).strftime("%d.%m.%Y")
        except ValueError: pass
    return make_certificate(u["full_name"], score, total, u["grade"], u["school"], date_str)

# ---------------------------------------------------------------- foydalanuvchi oqimi

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    upsert_user(uid, update.effective_user.username)
    context.user_data.clear()
    u = get_user(uid)
    if u["status"] == "in_progress":
        await update.message.reply_text(
            f"Testingiz hali tugamagan. Savol {len(json_load(u['answers']))+1}/{TOTAL_Q} dan davom etamiz.",
            reply_markup=main_keyboard(u))
        return
    if u["status"] == "completed":
        p = round((u["score"] or 0) / (u["total"] or TOTAL_Q) * 100)
        await update.message.reply_text(
            f"Assalomu alaykum, {u['full_name']}! Siz IQ testni topshirib bo‘lgansiz.\nNatijangiz: {p}%",
            reply_markup=main_keyboard(u))
        return
    await update.message.reply_text(
        "Assalomu alaykum! IQ Test botiga xush kelibsiz.\nQuyidan kerakli bo‘limni tanlang.",
        reply_markup=main_keyboard(u))

async def begin_registration(update, context, u):
    if u["status"] == "in_progress":
        await update.message.reply_text("Testingiz davom etmoqda. Avval testni tugating.", reply_markup=main_keyboard(u))
        return
    if u["status"] == "completed":
        await update.message.reply_text("Siz testni topshirib bo‘lgansiz.", reply_markup=main_keyboard(u))
        return
    set_fields(u["tg_id"], status="registration_name")
    await update.message.reply_text("Ism va familyangizni to‘liq kiriting:", reply_markup=back_keyboard())

def start_time_text():
    st = get_start_time()
    if not st:
        return "Test vaqti hali belgilanmagan. Vaqt e’lon qilinganda sizga xabar beramiz."
    if now() < st:
        return f"⏰ IQ test {fmt_dt(st)} da boshlanadi."
    return "✅ IQ test boshlandi! Quyidagi tugmani bosing."

async def start_test(update, context, u):
    uid = u["tg_id"]
    if u["status"] == "completed":
        await update.message.reply_text("Siz testni topshirib bo‘lgansiz.", reply_markup=main_keyboard(u)); return
    if u["status"] == "in_progress":
        await send_question(update.message, uid); return
    if u["status"] != "registered":
        await update.message.reply_text("Avval ro‘yxatdan o‘tishni tugating.", reply_markup=main_keyboard(u)); return
    st = get_start_time()
    if not st or now() < st:
        await update.message.reply_text(start_time_text(), reply_markup=main_keyboard(u)); return
    variant = random.choice(list(QUESTION_BANK))
    order = list(range(TOTAL_Q))
    random.shuffle(order)
    set_fields(uid, variant=variant, question_order=json.dumps(order), answers="[]",
               score=0, total=TOTAL_Q, status="in_progress", started_at=now_str(), finished_at=None)
    await update.message.reply_text(
        f"Test boshlandi! {TOTAL_Q} ta savolga javob bering. Javob tanlaganingizdan keyin keyingi savol chiqadi.",
        reply_markup=kb([[BTN_CONTINUE], [BTN_FINISH]]))
    await send_question(update.message, uid)

async def send_question(message, uid):
    u = get_user(uid)
    if not u or u["status"] != "in_progress":
        await message.reply_text("Faol testingiz topilmadi.", reply_markup=main_keyboard(u)); return
    order = json_load(u["question_order"])
    idx = len(json_load(u["answers"]))
    if idx >= len(order):
        await finish_test(message, uid); return
    q, opts, _ = QUESTION_BANK[u["variant"]][order[idx]]
    letters = "ABCD"
    buttons = [InlineKeyboardButton(f"{letters[i]}. {opts[i]}", callback_data=f"ans:{idx}:{i}") for i in range(4)]
    await message.reply_text(f"📝 Savol {idx+1}/{TOTAL_Q}\n\n{q}",
                             reply_markup=InlineKeyboardMarkup([buttons[:2], buttons[2:]]))

async def answer_callback(update, context):
    query = update.callback_query
    uid = query.from_user.id
    u = get_user(uid)
    if not u or u["status"] != "in_progress":
        await query.answer("Faol test topilmadi.", show_alert=True); return
    try:
        _, idxs, choices = query.data.split(":")
        idx, choice = int(idxs), int(choices)
    except ValueError:
        await query.answer(); return
    answers = json_load(u["answers"])
    if len(answers) != idx:
        await query.answer("Bu savolga javob berib bo‘lgansiz.", show_alert=True); return
    await query.answer()
    order = json_load(u["question_order"])
    qid = order[idx]
    q, opts, correct = QUESTION_BANK[u["variant"]][qid]
    answers.append({"question_id": qid, "choice": choice, "correct": choice == correct})
    set_fields(uid, answers=json.dumps(answers), score=sum(1 for a in answers if a["correct"]))
    try:
        await query.edit_message_text(f"📝 Savol {idx+1}/{TOTAL_Q}\n\n{q}\n\nJavobingiz: {'ABCD'[choice]}. {opts[choice]}")
    except Exception:
        pass
    if len(answers) >= TOTAL_Q:
        await query.message.reply_text("Oxirgi javob qabul qilindi. Sertifikatingiz tayyorlanmoqda…")
        await finish_test(query.message, uid, context)
    else:
        await send_question(query.message, uid)

async def finish_test(message, uid, context=None, early=False):
    u = get_user(uid)
    if not u or u["status"] != "in_progress":
        return
    answers = json_load(u["answers"])
    if len(answers) < TOTAL_Q and not early:
        return
    score = sum(1 for a in answers if a["correct"])
    set_fields(uid, status="completed", score=score, total=TOTAL_Q, finished_at=now_str())
    u = get_user(uid)
    percent = round(score / TOTAL_Q * 100)
    await message.reply_text(
        f"✅ Test yakunlandi!\n\nIsm-familya: {u['full_name']}\nMaktab: {u['school'] or '-'}\nSinf: {u['grade']}\n"
        f"Javob berilgan savollar: {len(answers)}/{TOTAL_Q}\nTo‘g‘ri javoblar: {score}/{TOTAL_Q}\nFoiz: {percent}%\n\nIzoh: {feedback_for(percent)}",
        reply_markup=main_keyboard(u))
    await message.reply_document(document=certificate_for_row(u), filename=safe_filename(u["full_name"], uid),
                                 caption="🏅 Al-Aziz Academy sertifikatingiz.")
    if context:
        await notify_admins(context.bot, f"✅ Test yakunlandi\n{u['full_name']} | {u['school'] or '-'} | {u['grade']}-sinf | {score}/{TOTAL_Q} ({percent}%)")

async def ask_finish(message, u):
    left = TOTAL_Q - len(json_load(u["answers"]))
    await message.reply_text(
        f"Testni yakunlamoqchimisiz?\nSiz hali {left} ta savolga javob bermadingiz. Ular noto‘g‘ri deb hisoblanadi va testni qayta ishlab bo‘lmaydi.",
        reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("✅ Ha, yakunlash", callback_data="fin:yes"),
                                            InlineKeyboardButton("↩️ Yo‘q, davom etish", callback_data="fin:no")]]))

async def finish_callback(update, context):
    query = update.callback_query
    uid = query.from_user.id
    u = get_user(uid)
    await query.answer()
    try: await query.edit_message_reply_markup(reply_markup=None)
    except Exception: pass
    if not u or u["status"] != "in_progress":
        return
    if query.data == "fin:yes":
        await query.message.reply_text("Test yakunlanmoqda. Sertifikatingiz tayyorlanmoqda…")
        await finish_test(query.message, uid, context, early=True)
    else:
        await send_question(query.message, uid)

async def notify_admins(bot, text):
    for aid in admin_ids():
        try: await bot.send_message(aid, text)
        except Exception: pass

async def broadcast(bot, text, statuses=("registered",)):
    marks = ",".join("?" * len(statuses))
    rows = sql(f"SELECT tg_id, status FROM users WHERE status IN ({marks})", list(statuses), "all")
    sent = 0
    for r in rows:
        try:
            await bot.send_message(r["tg_id"], text, reply_markup=main_keyboard(r))
            sent += 1
        except Exception:
            pass
        await asyncio.sleep(0.05)  # Telegram cheklovi: sekundiga ~30 xabar
    return sent

async def start_watcher(app):
    """Belgilangan vaqt kelganda ro‘yxatdan o‘tganlarga 'test boshlandi' xabarini yuboradi."""
    while True:
        try:
            st = get_start_time()
            key = get_setting("test_start")
            if st and now() >= st and get_setting("notified_start") != key:
                set_setting("notified_start", key)
                n = await broadcast(app.bot, f"🔔 IQ test boshlandi! «{BTN_START}» tugmasini bosing.")
                await notify_admins(app.bot, f"🔔 Test boshlandi. {n} ta foydalanuvchiga xabar yuborildi.")
        except Exception:
            log.exception("start_watcher xatosi")
        await asyncio.sleep(20)

async def post_init(app):
    app.bot_data["watcher"] = asyncio.create_task(start_watcher(app))

# ---------------------------------------------------------------- admin

async def send_long(message, text, **kw):
    chunks, cur = [], ""
    for line in text.split("\n"):
        if len(cur) + len(line) + 1 > 3800:
            chunks.append(cur); cur = ""
        cur += line + "\n"
    if cur.strip(): chunks.append(cur)
    for i, ch in enumerate(chunks):
        await message.reply_text(ch, **(kw if i == len(chunks) - 1 else {}))

STATUS_NAMES = {
    "in_progress": "⏳ Testni ishlayapti", "completed": "✅ Tugatgan",
    "registered": "🟢 Ro‘yxatdan o‘tgan, testni kutmoqda",
}

async def handle_admin(update, context, text):
    uid = update.effective_user.id
    msg = update.message
    if context.user_data.get("admin_step") == "set_time":
        context.user_data.pop("admin_step", None)
        try:
            st = datetime.strptime(text, "%d.%m.%Y %H:%M").replace(tzinfo=TZ)
        except ValueError:
            await msg.reply_text("❌ Format noto‘g‘ri. Masalan: 20.10.2026 10:00\nQaytadan «⏰ Test vaqtini belgilash» ni bosing.", reply_markup=admin_keyboard())
            return True
        if st <= now():
            await msg.reply_text("❌ Bu vaqt o‘tib ketgan. Kelajakdagi vaqtni kiriting.", reply_markup=admin_keyboard())
            return True
        set_setting("test_start", st.strftime("%Y-%m-%d %H:%M"))
        await msg.reply_text(f"✅ Test vaqti belgilandi: {fmt_dt(st)}.\nRo‘yxatdan o‘tganlarga xabar yuborilmoqda…", reply_markup=admin_keyboard())
        n = await broadcast(context.bot, f"⏰ IQ test {fmt_dt(st)} da boshlanadi. Shu vaqtda «{BTN_START}» tugmasini bosing.")
        await msg.reply_text(f"{n} ta foydalanuvchiga xabar yuborildi.")
        return True
    if text == BTN_SET_TIME:
        context.user_data["admin_step"] = "set_time"
        st = get_start_time()
        cur = f"Hozirgi vaqt: {fmt_dt(st)}\n\n" if st else ""
        await msg.reply_text(f"{cur}Test boshlanadigan kun va vaqtni kiriting (Toshkent vaqti).\nFormat: KUN.OY.YIL SOAT:DAQIQA\nMasalan: 20.10.2026 10:00", reply_markup=back_keyboard())
        return True
    if text == BTN_CLEAR_TIME:
        set_setting("test_start", None)
        await msg.reply_text("Test vaqti bekor qilindi. Endi hech kim testni boshlay olmaydi, toki yangi vaqt belgilanmaguncha.", reply_markup=admin_keyboard())
        return True
    if text == BTN_STATS:
        q = lambda query: sql(query.replace("COUNT(*)", "COUNT(*) AS n"), fetch="one")["n"]
        total = q("SELECT COUNT(*) FROM users")
        registered = q("SELECT COUNT(*) FROM users WHERE status IN ('registered','in_progress','completed')")
        progress = q("SELECT COUNT(*) FROM users WHERE status='in_progress'")
        done = q("SELECT COUNT(*) FROM users WHERE status='completed'")
        avg = sql("SELECT AVG(score*100.0/total) AS a FROM users WHERE status='completed'", fetch="one")["a"]
        st = get_start_time()
        await msg.reply_text(
            f"📊 BOT STATISTIKASI\n\nTest vaqti: {fmt_dt(st) if st else 'belgilanmagan'}\n\n"
            f"Botga kirganlar: {total}\nRo‘yxatdan o‘tganlar: {registered}\nTestni ishlayapti: {progress}\n"
            f"Testni tugatgan: {done}\nO‘rtacha natija: {round(avg) if avg is not None else '-'}%",
            reply_markup=admin_keyboard())
        return True
    if text == BTN_USERS:
        rows = sql("SELECT * FROM users WHERE full_name IS NOT NULL ORDER BY registered_at DESC", fetch="all")
        if not rows:
            await msg.reply_text("Hozircha foydalanuvchi yo‘q.", reply_markup=admin_keyboard()); return True
        lines = [f"👥 Foydalanuvchilar ({len(rows)} ta):"]
        for i, r in enumerate(rows, 1):
            status = STATUS_NAMES.get(r["status"], "📝 Ro‘yxatdan o‘tmoqda")
            res = ""
            if r["status"] == "completed":
                res = f" | {r['score']}/{r['total']} ({round(r['score']/r['total']*100)}%)"
            lines.append(f"\n{i}. {r['full_name']} | {r['phone'] or 'tel yo‘q'}\n   Maktab: {r['school'] or '-'} | Sinf: {r['grade'] or '-'}\n   {status}{res}")
        await send_long(msg, "\n".join(lines), reply_markup=admin_keyboard())
        return True
    if text == BTN_EXCEL:
        await msg.reply_document(document=export_excel(), filename=f"iq_test_natijalar_{now().strftime('%Y%m%d_%H%M')}.xlsx",
                                 caption="Barcha foydalanuvchilar va test natijalari.")
        return True
    if text == BTN_CERTS:
        rows = sql("SELECT * FROM users WHERE status='completed' ORDER BY finished_at", fetch="all")
        if not rows:
            await msg.reply_text("Hali hech kim testni tugatmagan.", reply_markup=admin_keyboard()); return True
        await msg.reply_text(f"{len(rows)} ta sertifikat tayyorlanmoqda…")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for r in rows:
                z.writestr(safe_filename(r["full_name"], r["tg_id"]), certificate_for_row(r).getvalue())
        buf.seek(0)
        await msg.reply_document(document=buf, filename=f"sertifikatlar_{now().strftime('%Y%m%d_%H%M')}.zip",
                                 caption=f"🏅 {len(rows)} ta sertifikat.", reply_markup=admin_keyboard())
        return True
    if text == BTN_LOGOUT:
        sql("DELETE FROM admins WHERE tg_id=?", (uid,))
        await msg.reply_text("Admin paneldan chiqdingiz.", reply_markup=main_keyboard(get_user(uid)))
        return True
    return False

EXCEL_STATUS = {"in_progress": "Testni ishlayapti", "completed": "Tugatgan", "registered": "Testni kutmoqda"}

def export_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Foydalanuvchilar"
    ws.append(["№", "Telegram ID", "Username", "Ism-familya", "Telefon", "Maktab", "Sinf", "Variant", "Holat",
               "Ball", "Jami savol", "Foiz", "Izoh", "Test boshlagan", "Test tugatgan", "Ro‘yxatdan o‘tgan"])
    rows = sql("SELECT * FROM users ORDER BY registered_at", fetch="all")
    for i, r in enumerate(rows, 1):
        total = r["total"] or TOTAL_Q
        done = r["status"] == "completed"
        percent = round((r["score"] or 0) / total * 100) if done else ""
        ws.append([i, r["tg_id"], r["username"], r["full_name"], r["phone"], r["school"], r["grade"], r["variant"],
                   EXCEL_STATUS.get(r["status"], "Ro‘yxatdan o‘tmoqda"),
                   r["score"] if done else "", total if done else "", percent,
                   feedback_for(percent) if done else "", r["started_at"], r["finished_at"], r["registered_at"]])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="13315C")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 45)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

# ---------------------------------------------------------------- matnli xabarlar

async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message: return
    msg = update.message
    text = (msg.text or "").strip()
    uid = update.effective_user.id
    upsert_user(uid, update.effective_user.username)
    u = get_user(uid)

    if text == BTN_MENU:
        context.user_data.clear()
        if u["status"] in ("registration_name", "registration_phone", "registration_school", "registration_grade"):
            set_fields(uid, status="new" if not u["grade"] else "registered")
            u = get_user(uid)
        await msg.reply_text("Bosh menyu:", reply_markup=main_keyboard(u))
        return

    # Admin kirishi
    step = context.user_data.get("admin_step")
    if text == BTN_ADMIN:
        if is_admin(uid):
            await msg.reply_text("Admin panel:", reply_markup=admin_keyboard()); return
        context.user_data["admin_step"] = "login"
        await msg.reply_text("Admin loginini kiriting:", reply_markup=back_keyboard()); return
    if step == "login":
        context.user_data["admin_login_candidate"] = text
        context.user_data["admin_step"] = "password"
        await msg.reply_text("Admin parolini kiriting:"); return
    if step == "password":
        login = context.user_data.pop("admin_login_candidate", "")
        context.user_data.pop("admin_step", None)
        try: await msg.delete()  # parol chatda qolmasin
        except Exception: pass
        if login == ADMIN_LOGIN and ADMIN_PASSWORD and text == ADMIN_PASSWORD:
            sql("INSERT INTO admins(tg_id) VALUES(?) ON CONFLICT(tg_id) DO NOTHING", (uid,))
            await msg.reply_text("✅ Admin panelga muvaffaqiyatli kirdingiz.", reply_markup=admin_keyboard())
        else:
            await asyncio.sleep(2)
            await msg.reply_text("❌ Login yoki parol noto‘g‘ri.", reply_markup=main_keyboard(u))
        return
    if is_admin(uid) and await handle_admin(update, context, text):
        return

    # Ro‘yxatdan o‘tish
    status = u["status"]
    if text == BTN_REGISTER:
        await begin_registration(update, context, u); return
    if text == BTN_START:
        await start_test(update, context, u); return
    if text == BTN_FINISH:
        if status == "in_progress": await ask_finish(msg, u)
        else: await msg.reply_text("Sizda faol test yo‘q.", reply_markup=main_keyboard(u))
        return
    if text == BTN_CONTINUE:
        await send_question(msg, uid); return

    if status == "registration_name":
        if len(text.split()) < 2 or len(text) < 4:
            await msg.reply_text("Iltimos, ism va familyani to‘liq kiriting. Masalan: Ali Valiyev"); return
        set_fields(uid, full_name=text[:100], status="registration_phone")
        await msg.reply_text("Telefon raqamingizni tugma orqali yuboring:", reply_markup=phone_keyboard()); return
    if status == "registration_phone":
        await msg.reply_text("Telefon raqamni pastdagi «📱 Telefon raqamni yuborish» tugmasi orqali yuboring.", reply_markup=phone_keyboard()); return
    if status == "registration_school":
        if len(text) < 2:
            await msg.reply_text("Maktabingiz nomi yoki raqamini kiriting. Masalan: 12-maktab"); return
        set_fields(uid, school=text[:100], status="registration_grade")
        await msg.reply_text("Sinfingizni tanlang (1–11):", reply_markup=grade_keyboard()); return
    if status == "registration_grade":
        if not text.isdigit() or not 1 <= int(text) <= 11:
            await msg.reply_text("1 dan 11 gacha bo‘lgan sinfni tugma orqali tanlang.", reply_markup=grade_keyboard()); return
        set_fields(uid, grade=int(text), status="registered")
        u = get_user(uid)
        await msg.reply_text(f"✅ Ro‘yxatdan o‘tdingiz!\n\n{start_time_text()}", reply_markup=main_keyboard(u)); return

    await msg.reply_text("Iltimos, menyudagi tugmalardan foydalaning.", reply_markup=main_keyboard(u))

async def contact_handler(update, context):
    uid = update.effective_user.id
    u = get_user(uid)
    if not u or u["status"] != "registration_phone":
        await update.message.reply_text("Iltimos, menyudagi tugmalardan foydalaning.", reply_markup=main_keyboard(u)); return
    contact = update.message.contact
    if not contact.user_id or contact.user_id != uid:
        await update.message.reply_text("Iltimos, o‘zingizning telefon raqamingizni tugma orqali yuboring.", reply_markup=phone_keyboard()); return
    set_fields(uid, phone=contact.phone_number, status="registration_school")
    await update.message.reply_text("Maktabingizni kiriting (masalan: 12-maktab):", reply_markup=back_keyboard())

async def error_handler(update, context):
    log.exception("Unhandled exception", exc_info=context.error)
    try:
        if isinstance(update, Update) and update.effective_message:
            await update.effective_message.reply_text("Xatolik yuz berdi. Qayta urinib ko‘ring yoki administratorga murojaat qiling.")
    except Exception:
        pass

def main():
    if not BOT_TOKEN:
        raise SystemExit("BOT_TOKEN .env faylida ko‘rsatilmagan. README faylini ko‘ring.")
    if not ADMIN_PASSWORD:
        raise SystemExit("ADMIN_PASSWORD .env faylida ko‘rsatilmagan. README faylini ko‘ring.")
    init_db()
    log.info("Baza: %s", "PostgreSQL" if USE_PG else f"SQLite ({DB_PATH})")
    app = Application.builder().token(BOT_TOKEN).post_init(post_init).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CallbackQueryHandler(answer_callback, pattern=r"^ans:"))
    app.add_handler(CallbackQueryHandler(finish_callback, pattern=r"^fin:"))
    app.add_handler(MessageHandler(filters.CONTACT, contact_handler))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_error_handler(error_handler)
    log.info("IQ test bot ishga tushdi.")
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == "__main__":
    main()
